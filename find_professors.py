import os
import re
import time
import html
import unicodedata
import requests
import pandas as pd
from openpyxl import load_workbook

from datetime import datetime, timezone, timedelta
from difflib import SequenceMatcher

# Optional: used for stronger text-similarity scoring.
try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    HAS_SKLEARN = True
except ImportError:
    HAS_SKLEARN = False


# ==============================================================================
# CONFIG
# ==============================================================================

TELEGRAM_BOT_TOKEN = os.environ.get("TELEGRAM_BOT_TOKEN")
TELEGRAM_CHAT_ID = os.environ.get("TELEGRAM_CHAT_ID")
OPENALEX_API_KEY = os.environ.get("OPENALEX_API_KEY")  # Optional.

# Detect whichever tracking file is present in the repository.
if os.path.exists("List of emails.xlsx"):
    FILE_NAME = "List of emails.xlsx"
elif os.path.exists("List of emails.xlsx - Sheet1.csv"):
    FILE_NAME = "List of emails.xlsx - Sheet1.csv"
else:
    FILE_NAME = "List of emails.xlsx"

BATCH_SIZE = 5
MAX_FACULTY_PER_UNIVERSITY = 5

# Recent publication window.
RECENT_YEARS = 5

# Search breadth.
WORKS_PER_GROUP = 15
MAX_CANDIDATE_AUTHORS = 20
MAX_WORKS_FOR_FIT = 10

# Optional homepage crawl for funding evidence.
ENABLE_FUNDING_PAGE_CHECK = True
FUNDING_PAGE_TIMEOUT = 12
MAX_HOMEPAGE_CHARS = 50000

OPENALEX_BASE = "https://api.openalex.org"
OPENALEX_EMAIL = "phd-lead-finder@academic.org"

# ==============================================================================
# CANDIDATE RESEARCH PROFILE
# ==============================================================================

# Strong terms receive much larger weights than generic AI/ML/robotics terms.
# This is deliberately asymmetric: a professor doing VIO should outrank
# someone doing generic machine learning.
SEARCH_GROUPS = [
    (
        "Visual-Inertial Navigation & SLAM",
        [
            "visual inertial odometry",
            "visual inertial navigation",
            "visual SLAM",
            "visual-inertial SLAM",
            "visual odometry",
            "camera IMU fusion",
            "inertial navigation",
        ],
    ),
    (
        "Robotics & Autonomous Navigation",
        [
            "robot localization",
            "autonomous navigation",
            "mobile robot navigation",
            "robot perception",
            "autonomous robotics",
            "field robotics",
            "outdoor robotics",
        ],
    ),
    (
        "GNSS & Multi-Sensor Navigation",
        [
            "GNSS positioning",
            "GNSS navigation",
            "GNSS INS integration",
            "GNSS sensor fusion",
            "multi-sensor fusion",
            "robust localization",
            "precise positioning",
        ],
    ),
    (
        "Computer Vision & 3D Perception",
        [
            "computer vision robotics",
            "3D computer vision",
            "visual perception",
            "3D perception",
            "visual localization",
            "depth estimation",
            "stereo vision",
        ],
    ),
    (
        "LiDAR & Multimodal Sensor Fusion",
        [
            "LiDAR camera fusion",
            "LiDAR inertial fusion",
            "LiDAR visual odometry",
            "LiDAR SLAM",
            "multimodal sensor fusion",
            "camera LiDAR IMU",
        ],
    ),
    (
        "State Estimation & Optimization",
        [
            "state estimation robotics",
            "pose estimation",
            "trajectory estimation",
            "factor graph optimization",
            "nonlinear optimization robotics",
            "probabilistic robotics",
            "Kalman filtering robotics",
        ],
    ),
    (
        "Machine Learning for Robotics & Perception",
        [
            "deep learning robotics",
            "machine learning robotics",
            "deep learning computer vision",
            "self-supervised visual navigation",
            "learning-based localization",
            "neural visual odometry",
            "learning-based perception",
        ],
    ),
    (
        "Autonomous Vehicles & Intelligent Systems",
        [
            "autonomous vehicles",
            "autonomous driving perception",
            "vehicle localization",
            "intelligent transportation systems",
            "robotic navigation",
            "autonomous systems perception",
        ],
    ),
    (
        "Embedded AI & Edge Robotics",
        [
            "embedded AI robotics",
            "edge AI robotics",
            "real-time computer vision",
            "GPU accelerated robotics",
            "embedded computer vision",
            "AI acceleration robotics",
            "Jetson robotics",
        ],
    ),
    (
        "Medical Image Computing",
        [
            "medical image segmentation",
            "medical image processing",
            "deep learning medical imaging",
            "computer vision medical imaging",
            "image segmentation",
        ],
    ),
    (
        "Embedded Systems & Digital Hardware",
        [
            "embedded systems",
            "embedded AI",
            "FPGA neural network accelerator",
            "hardware acceleration",
            "digital hardware design",
            "real-time embedded systems",
        ],
    ),
]

# Explicitly weighted research concepts.
# Lower-weight generic terms prevent "AI", "robotics", etc. from dominating.
TIER1_PHRASES = {
    "visual inertial odometry": 10.0,
    "visual inertial navigation": 10.0,
    "visual-inertial slam": 10.0,
    "visual slam": 9.5,
    "visual odometry": 9.0,
    "camera imu fusion": 9.0,
    "gnss ins integration": 9.0,
    "gnss sensor fusion": 9.0,
    "gnss positioning": 8.5,
    "robot localization": 8.5,
    "visual localization": 8.5,
    "autonomous navigation": 8.5,
    "factor graph optimization": 8.0,
    "state estimation": 8.0,
    "lidar slam": 8.0,
    "lidar camera fusion": 8.0,
    "camera lidar imu": 8.0,
    "inertial navigation": 8.0,
    "multi-sensor fusion": 8.0,
    "sensor fusion": 7.5,
    "robust localization": 7.5,
    "pose estimation": 7.0,
    "trajectory estimation": 7.0,
    "computer vision robotics": 7.0,
    "robot perception": 7.0,
}

TIER2_PHRASES = {
    "3d computer vision": 5.5,
    "3d perception": 5.5,
    "visual perception": 5.5,
    "stereo vision": 5.0,
    "depth estimation": 5.0,
    "field robotics": 5.0,
    "mobile robot navigation": 5.0,
    "autonomous vehicles": 5.0,
    "autonomous driving perception": 5.0,
    "probabilistic robotics": 5.0,
    "neural visual odometry": 5.0,
    "learning-based localization": 4.5,
    "self-supervised visual navigation": 4.5,
    "learning-based perception": 4.5,
    "lidar inertial fusion": 5.5,
    "multimodal sensor fusion": 5.5,
}

TIER3_PHRASES = {
    "deep learning": 2.5,
    "machine learning": 2.0,
    "computer vision": 2.0,
    "object detection": 2.0,
    "object tracking": 2.0,
    "embedded ai": 2.0,
    "edge ai": 2.0,
    "cuda": 1.5,
    "gpu": 1.5,
    "ros": 1.5,
    "ros2": 1.5,
    "embedded systems": 1.5,
    "fpga": 1.0,
}

RESEARCH_PROFILE_TEXT = " ".join(
    phrase for _, phrases in SEARCH_GROUPS for phrase in phrases
)

# Funding phrases only count as evidence when found on a faculty/homepage page.
FUNDING_HIGH = [
    "fully funded phd",
    "funded phd",
    "funded doctoral",
    "phd studentship",
    "doctoral studentship",
    "phd scholarship",
    "funded position",
    "funded positions",
    "phd opportunity",
    "phd opportunities",
]

FUNDING_MEDIUM = [
    "phd positions",
    "doctoral positions",
    "phd students",
    "doctoral students",
    "research assistantship",
    "research assistant",
    "scholarship",
    "studentship",
]

# ==============================================================================
# HTTP HELPERS
# ==============================================================================

session = requests.Session()
session.headers.update(
    {
        "User-Agent": f"Mozilla/5.0 PhDLeadFinder/2.0 (+mailto:{OPENALEX_EMAIL})",
        "Accept-Language": "en-US,en;q=0.9",
    }
)


def openalex_get(path_or_url, params=None, timeout=20, retries=3):
    """GET an OpenAlex endpoint with retries and useful diagnostics."""
    if path_or_url.startswith("http://") or path_or_url.startswith("https://"):
        url = path_or_url
    else:
        url = f"{OPENALEX_BASE}/{path_or_url.lstrip('/')}"

    params = dict(params or {})
    params.setdefault("mailto", OPENALEX_EMAIL)

    if OPENALEX_API_KEY:
        params.setdefault("api_key", OPENALEX_API_KEY)

    last_error = None

    for attempt in range(retries):
        try:
            response = session.get(url, params=params, timeout=timeout)
            if response.status_code == 200:
                return response.json()

            if response.status_code in (429, 500, 502, 503, 504):
                wait = min(10 * (attempt + 1), 30)
                print(f"    [!] OpenAlex HTTP {response.status_code}; retrying in {wait}s...")
                time.sleep(wait)
                continue

            print(f"    [!] OpenAlex HTTP {response.status_code}: {response.text[:250]}")
            return None

        except Exception as exc:
            last_error = exc
            time.sleep(2 * (attempt + 1))

    print(f"    [!] OpenAlex request failed after retries: {last_error}")
    return None


# ==============================================================================
# TEXT / NORMALIZATION
# ==============================================================================

def normalize_name(value):
    value = str(value or "").lower()
    value = re.sub(r"[^\w\s]", " ", value, flags=re.UNICODE)
    return re.sub(r"\s+", " ", value).strip()


def clean_university_name(name):
    """
    Clean Excel annotations without incorrectly splitting legitimate names such as
    'University of Science and Technology of China'.
    """
    name = str(name or "").strip()
    name = re.sub(r"\s+", " ", name)
    name = re.sub(r"\([^)]*\)", "", name).strip()

    # Remove common spreadsheet notes after separators.
    name = re.sub(r"\s*-\s*3years?\s+phd\s+programs?.*$", "", name, flags=re.I)
    name = re.sub(r"\s*\(\s*3years?\s+phd programs?\s*\)\s*$", "", name, flags=re.I)

    # Known annotation in the user's sheet.
    if normalize_name(name) == normalize_name("ETH Zürich and Bologna"):
        name = "ETH Zürich"

    return name.strip(" -")


# Canonical aliases used ONLY for duplicate detection / batch selection.
# The original Excel text is preserved in the University column.
UNIVERSITY_ALIASES = {
    "eth zurich": "eth zurich",
    "eth zürich": "eth zurich",
    "eth zurich and bologna": "eth zurich",
    "technical university of munich": "technical university of munich",
    "technical university of munich tum": "technical university of munich",
    "tu munich": "technical university of munich",
    "rwth aachen": "rwth aachen university",
    "rwth aachen university": "rwth aachen university",
    "delft university of technology": "delft university of technology",
    "tu delft": "delft university of technology",
    "eindhoven university of technology": "eindhoven university of technology",
    "tu eindhoven": "eindhoven university of technology",
    "university of twente": "university of twente",
    "radbound university": "radboud university",
    "radboud university": "radboud university",
    "national university of singapore nus": "national university of singapore",
    "national university of singapore": "national university of singapore",
    "nanyang technological university singapore ntu singapore": "nanyang technological university singapore",
    "nanyang technological university singapore": "nanyang technological university singapore",
}


def university_key(name):
    """Return a stable key so duplicates are processed once per run."""
    cleaned = clean_university_name(name)
    normalized = unicodedata.normalize("NFKD", cleaned).encode("ascii", "ignore").decode("ascii")
    normalized = normalize_name(normalized)
    normalized = re.sub(r"^the\s+", "", normalized)
    normalized = re.sub(r"\b(university|universitaet|university)\b\s*$", "university", normalized)
    return UNIVERSITY_ALIASES.get(normalized, normalized)


def duplicate_indices_for_university(df, key):
    """Return every row representing the same university key."""
    return [idx for idx in df.index if university_key(df.at[idx, "University"]) == key]


def reconstruct_abstract(inverted_index):
    if not inverted_index:
        return ""

    positions = []
    for word, indexes in inverted_index.items():
        for idx in indexes:
            positions.append((idx, word))

    positions.sort(key=lambda x: x[0])
    return " ".join(word for _, word in positions)


def work_text(work):
    parts = [
        work.get("title", ""),
        reconstruct_abstract(work.get("abstract_inverted_index")),
    ]

    for topic in work.get("topics") or []:
        topic_name = topic.get("display_name", "")
        if topic_name:
            parts.append(topic_name)

    for keyword in work.get("keywords") or []:
        keyword_name = keyword.get("display_name", "")
        if keyword_name:
            parts.append(keyword_name)

    return " ".join(parts).strip()


def phrase_score(text):
    """
    Weighted phrase matching.
    Tier-1 navigation concepts dominate generic AI/ML keywords.
    """
    lower = normalize_name(text)
    total = 0.0
    matched = []

    all_phrases = [
        ("T1", TIER1_PHRASES),
        ("T2", TIER2_PHRASES),
        ("T3", TIER3_PHRASES),
    ]

    for tier, mapping in all_phrases:
        for phrase, weight in mapping.items():
            normalized_phrase = normalize_name(phrase)
            if normalized_phrase in lower:
                total += weight
                matched.append((phrase, weight, tier))

    # Mild bonus for explicitly strong combinations.
    combo_pairs = [
        ("gnss", "sensor fusion"),
        ("camera", "imu"),
        ("lidar", "slam"),
        ("visual", "navigation"),
        ("robot", "localization"),
        ("factor graph", "navigation"),
    ]
    for a, b in combo_pairs:
        if a in lower and b in lower:
            total += 2.0

    matched.sort(key=lambda x: x[1], reverse=True)
    return total, matched


def tfidf_similarity(profile_text, texts):
    """
    TF-IDF similarity is used as a lightweight semantic-text similarity layer.
    It is optional; if scikit-learn is unavailable, the algorithm falls back
    to weighted phrase matching.
    """
    if not texts:
        return []

    if not HAS_SKLEARN:
        return [0.0 for _ in texts]

    try:
        vectorizer = TfidfVectorizer(
            stop_words="english",
            ngram_range=(1, 2),
            min_df=1,
            max_features=6000,
        )
        matrix = vectorizer.fit_transform([profile_text] + texts)
        sims = cosine_similarity(matrix[0:1], matrix[1:]).flatten()
        return [float(max(0.0, s)) for s in sims]
    except Exception:
        return [0.0 for _ in texts]


def recent_year_start():
    dt = datetime.now(timezone.utc) - timedelta(days=365.25 * RECENT_YEARS)
    return dt.date().isoformat()


def recency_factor(publication_date):
    if not publication_date:
        return 0.2

    try:
        dt = datetime.fromisoformat(publication_date)
    except ValueError:
        return 0.2

    now = datetime.now(timezone.utc)
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)

    days = max(0, (now - dt).days)

    if days <= 180:
        return 1.00
    if days <= 365:
        return 0.90
    if days <= 730:
        return 0.78
    if days <= 1095:
        return 0.68
    if days <= 1460:
        return 0.58
    return 0.45


# ==============================================================================
# INSTITUTION RESOLUTION
# ==============================================================================

def institution_name_similarity(query_name, candidate_name):
    q = normalize_name(query_name)
    c = normalize_name(candidate_name)

    q_tokens = set(q.split())
    c_tokens = set(c.split())

    token_score = len(q_tokens & c_tokens) / max(len(q_tokens), 1)
    sequence_score = SequenceMatcher(None, q, c).ratio()

    return 0.65 * token_score + 0.35 * sequence_score


def resolve_institution(raw_uni_name):
    clean_name = clean_university_name(raw_uni_name)
    print(f"--> Resolving OpenAlex institution: {clean_name}")

    data = openalex_get(
        "/institutions",
        params={
            "search": clean_name,
            "per_page": 10,
        },
    )

    if not data or not data.get("results"):
        print(f"    [!] Institution not resolved: {clean_name}")
        return None

    ranked = sorted(
        data["results"],
        key=lambda item: institution_name_similarity(
            clean_name, item.get("display_name", "")
        ),
        reverse=True,
    )

    best = ranked[0]
    similarity = institution_name_similarity(
        clean_name, best.get("display_name", "")
    )

    if similarity < 0.45:
        print(
            f"    [!] Weak institution match: {best.get('display_name')} "
            f"(similarity={similarity:.2f})"
        )
        return None

    print(
        f"    [✓] Matched: {best.get('display_name')} "
        f"({best.get('id')}, similarity={similarity:.2f})"
    )
    return best


# ==============================================================================
# WORK SEARCH + FACULTY AGGREGATION
# ==============================================================================

def build_group_query(keywords):
    return " OR ".join(f'"{keyword}"' for keyword in keywords)


def author_belongs_to_institution(authorship, institution_id):
    target_id = institution_id.rstrip("/").split("/")[-1]
    for inst in authorship.get("institutions") or []:
        inst_id = (inst.get("id") or "").rstrip("/").split("/")[-1]
        if inst_id == target_id:
            return True
    return False


def search_faculty_broad(raw_uni_name):
    """
    Returns:
      {
        "ok": bool,
        "institution": dict,
        "faculty": [faculty_record, ...]
      }
    """
    institution = resolve_institution(raw_uni_name)

    if not institution:
        return {"ok": False, "institution": None, "faculty": []}

    inst_id = institution["id"].rstrip("/").split("/")[-1]
    start_date = recent_year_start()

    all_works = {}
    author_stats = {}

    for field_label, keywords in SEARCH_GROUPS:
        query_str = build_group_query(keywords)

        data = openalex_get(
            "/works",
            params={
                "filter": (
                    f"authorships.institutions.id:{inst_id},"
                    f"from_publication_date:{start_date}"
                ),
                "search": query_str,
                "sort": "relevance_score:desc,publication_date:desc",
                "per_page": WORKS_PER_GROUP,
                # Keep payload reasonably small.
                "select": (
                    "id,title,publication_date,publication_year,"
                    "authorships,abstract_inverted_index,topics,keywords,"
                    "cited_by_count"
                ),
            },
        )

        if not data:
            continue

        for work in data.get("results", []):
            work_id = work.get("id")
            if not work_id:
                continue

            # Keep the work once, but remember every matched group.
            if work_id not in all_works:
                all_works[work_id] = {
                    "work": work,
                    "fields": set(),
                }

            all_works[work_id]["fields"].add(field_label)

    if not all_works:
        return {"ok": True, "institution": institution, "faculty": []}

    work_records = list(all_works.values())
    texts = [work_text(item["work"]) for item in work_records]
    sims = tfidf_similarity(RESEARCH_PROFILE_TEXT, texts)

    for idx, item in enumerate(work_records):
        work = item["work"]
        text = texts[idx]

        weighted_score, matched = phrase_score(text)
        # Normalize weighted phrase score into 0..1.
        # 45 is a deliberately high "excellent direct fit" reference point.
        keyword_norm = min(weighted_score / 45.0, 1.0)

        # Optional TF-IDF similarity. It is a supporting signal, never the
        # dominant signal, so generic ML papers do not rise to the top.
        similarity_norm = min(sims[idx] * 3.0, 1.0)

        base_fit = 0.80 * keyword_norm + 0.20 * similarity_norm
        recency = recency_factor(work.get("publication_date"))

        # More recent matching papers count more.
        contribution = base_fit * (0.70 + 0.30 * recency)

        for authorship in work.get("authorships") or []:
            if not author_belongs_to_institution(authorship, inst_id):
                continue

            author = authorship.get("author") or {}
            author_id = author.get("id")
            author_name = author.get("display_name")

            if not author_id or not author_name:
                continue

            if author_id not in author_stats:
                author_stats[author_id] = {
                    "author_id": author_id,
                    "professor": author_name,
                    "matched_works": {},
                    "fields": {},
                }

            author_stats[author_id]["matched_works"][work.get("id")] = {
                "title": work.get("title", "N/A"),
                "publication_date": work.get("publication_date", ""),
                "year": work.get("publication_year", ""),
                "cited_by_count": work.get("cited_by_count", 0),
                "fit": contribution,
                "raw_fit": base_fit,
                "matched_phrases": [x[0] for x in matched[:8]],
            }

            for field in item["fields"]:
                author_stats[author_id]["fields"][field] = (
                    author_stats[author_id]["fields"].get(field, 0) + contribution
                )

    # Preliminary ranking: strong research fit + breadth.
    for author_id, rec in author_stats.items():
        works = list(rec["matched_works"].values())

        if not works:
            rec["pre_score"] = 0.0
            continue

        total_fit = sum(w["fit"] for w in works)
        best_fit = max(w["raw_fit"] for w in works)
        breadth = len(rec["fields"])

        rec["pre_score"] = (
            min(total_fit / 5.0, 5.0)
            + min(best_fit * 4.0, 4.0)
            + min(breadth * 0.25, 1.0)
        )

    top_candidates = sorted(
        author_stats.values(),
        key=lambda x: x["pre_score"],
        reverse=True,
    )[:MAX_CANDIDATE_AUTHORS]

    # Enrich the strongest candidate authors.
    faculty = []
    for rec in top_candidates:
        author_data = get_author_details(rec["author_id"])
        if not author_data:
            author_data = {}

        recent_author_works = get_recent_author_works(
            rec["author_id"],
            recent_year_start(),
            per_page=25,
        )

        enriched = score_faculty_record(
            raw_uni_name=raw_uni_name,
            institution=institution,
            candidate=rec,
            author_data=author_data,
            recent_author_works=recent_author_works,
        )
        faculty.append(enriched)

        time.sleep(0.15)

    faculty.sort(
        key=lambda x: x["faculty_score"],
        reverse=True,
    )

    return {
        "ok": True,
        "institution": institution,
        "faculty": faculty[:MAX_FACULTY_PER_UNIVERSITY],
    }


def get_author_details(author_id):
    return openalex_get(
        f"/authors/{author_id.rstrip('/').split('/')[-1]}",
        timeout=15,
    )


def get_recent_author_works(author_id, start_date, per_page=25):
    data = openalex_get(
        "/works",
        params={
            "filter": (
                f"author.id:{author_id.rstrip('/').split('/')[-1]},"
                f"from_publication_date:{start_date}"
            ),
            "sort": "publication_date:desc",
            "per_page": per_page,
            "select": (
                "id,title,publication_date,publication_year,cited_by_count,"
                "authorships"
            ),
        },
        timeout=18,
    )
    return (data or {}).get("results", [])


# ==============================================================================
# FACULTY SCORING
# ==============================================================================

def activity_score(recent_works):
    count = len(recent_works)

    if count >= 10:
        count_score = 10.0
    elif count >= 7:
        count_score = 9.0
    elif count >= 5:
        count_score = 8.0
    elif count >= 3:
        count_score = 6.5
    elif count >= 2:
        count_score = 5.0
    elif count == 1:
        count_score = 3.0
    else:
        count_score = 0.0

    latest_date = ""
    if recent_works:
        latest_date = max(
            [w.get("publication_date", "") for w in recent_works if w.get("publication_date", "")],
            default="",
        )

    latest_factor = recency_factor(latest_date)
    latest_score = latest_factor * 10.0

    return round(0.65 * count_score + 0.35 * latest_score, 2), latest_date


def detect_funding_from_homepage(homepage_url):
    if not ENABLE_FUNDING_PAGE_CHECK or not homepage_url:
        return {
            "score": 0.0,
            "evidence": "No homepage/funding evidence available",
            "url": homepage_url or "",
        }

    try:
        response = session.get(
            homepage_url,
            timeout=FUNDING_PAGE_TIMEOUT,
            allow_redirects=True,
        )
        if response.status_code != 200:
            return {
                "score": 0.0,
                "evidence": f"Homepage unavailable (HTTP {response.status_code})",
                "url": homepage_url,
            }

        text = response.text[:MAX_HOMEPAGE_CHARS].lower()
        text = re.sub(r"\s+", " ", text)

        high_hits = [p for p in FUNDING_HIGH if p in text]
        medium_hits = [p for p in FUNDING_MEDIUM if p in text]

        if high_hits:
            return {
                "score": 10.0,
                "evidence": "Funding-related language: " + ", ".join(high_hits[:4]),
                "url": response.url,
            }

        if medium_hits:
            return {
                "score": 6.0,
                "evidence": "PhD/research-student language: " + ", ".join(medium_hits[:4]),
                "url": response.url,
            }

        return {
            "score": 0.0,
            "evidence": "No explicit PhD funding evidence found on homepage",
            "url": response.url,
        }

    except Exception as exc:
        return {
            "score": 0.0,
            "evidence": f"Funding page check failed: {type(exc).__name__}",
            "url": homepage_url,
        }


def score_faculty_record(
    raw_uni_name,
    institution,
    candidate,
    author_data,
    recent_author_works,
):
    matched_works = list(candidate["matched_works"].values())
    matched_works.sort(
        key=lambda w: (
            w.get("fit", 0.0),
            w.get("publication_date", ""),
            w.get("cited_by_count", 0),
        ),
        reverse=True,
    )

    top_fit_values = [w["raw_fit"] for w in matched_works[:MAX_WORKS_FOR_FIT]]

    if top_fit_values:
        # Best paper matters, but repeated high-quality matches matter too.
        best = max(top_fit_values)
        mean_top = sum(top_fit_values) / len(top_fit_values)
        research_score = min((0.60 * best + 0.40 * mean_top) * 10.0, 10.0)
    else:
        research_score = 0.0

    activity, latest_publication = activity_score(recent_author_works)

    homepage_url = author_data.get("homepage_url") or ""

    funding = detect_funding_from_homepage(homepage_url)

    faculty_score = (
        0.60 * research_score
        + 0.20 * activity
        + 0.20 * funding["score"]
    )

    field_scores = candidate["fields"]
    top_fields = [
        field
        for field, _score in sorted(
            field_scores.items(),
            key=lambda x: x[1],
            reverse=True,
        )[:3]
    ]

    top_work = matched_works[0] if matched_works else {}
    recent_title = top_work.get("title", "N/A")
    recent_year = top_work.get("year", "")

    return {
        "field": top_fields[0] if top_fields else "General CV/Robotics",
        "fields": top_fields,
        "professor": candidate["professor"],
        "university": raw_uni_name,
        "author_id": candidate["author_id"],
        "profile_url": author_data.get("id", candidate["author_id"]),
        "homepage_url": homepage_url,
        "paper": recent_title,
        "year": recent_year,
        "research_fit_score": round(research_score, 2),
        "activity_score": round(activity, 2),
        "funding_score": round(funding["score"], 2),
        "faculty_score": round(faculty_score, 2),
        "matched_works_count": len(matched_works),
        "recent_works_count": len(recent_author_works),
        "latest_publication": latest_publication,
        "funding_evidence": funding["evidence"],
        "funding_url": funding["url"],
        "recent_works": matched_works[:3],
    }


# ==============================================================================
# TELEGRAM
# ==============================================================================

def send_telegram_html(text):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        print("[!] Missing Telegram credentials.")
        return False

    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage"
    payload = {
        "chat_id": TELEGRAM_CHAT_ID,
        "text": text,
        "parse_mode": "HTML",
        "disable_web_page_preview": True,
    }

    try:
        response = requests.post(url, json=payload, timeout=20)

        if response.status_code != 200:
            print(f"[!] Telegram HTTP {response.status_code}: {response.text[:500]}")
            return False

        body = response.json()
        if not body.get("ok"):
            print(f"[!] Telegram API error: {body}")
            return False

        return True

    except Exception as exc:
        print(f"[!] Telegram send error: {exc}")
        return False


def send_telegram_document(file_path, caption=""):
    if not TELEGRAM_BOT_TOKEN or not TELEGRAM_CHAT_ID:
        return False

    url = f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendDocument"

    try:
        with open(file_path, "rb") as doc:
            response = requests.post(
                url,
                data={
                    "chat_id": TELEGRAM_CHAT_ID,
                    "caption": caption,
                },
                files={"document": doc},
                timeout=35,
            )

        if response.status_code != 200:
            print(f"[!] Telegram document HTTP {response.status_code}: {response.text[:500]}")
            return False

        body = response.json()
        if not body.get("ok"):
            print(f"[!] Telegram document API error: {body}")
            return False

        return True

    except Exception as exc:
        print(f"[!] Document send error: {exc}")
        return False


# ==============================================================================
# SHEET OUTPUT
# ==============================================================================

OUTPUT_COLUMNS = [
    "Checked_General?",
    "Discovered_Field",
    "General_PI",
    "General_Paper",
    "University_Research_Fit",
    "University_Activity",
    "University_Funding_Evidence",
    "University_Score",
    "Faculty_Breadth",
    "Top_Faculty",
    "Top_Faculty_Details",
    "Funding_Evidence_Details",
    "Last_Publication",
    "Note",
]


def ensure_output_columns(df):
    for col in OUTPUT_COLUMNS:
        if col not in df.columns:
            if col == "Checked_General?":
                df[col] = "NO"
            else:
                df[col] = ""
    return df


def format_faculty_details(faculty):
    if not faculty:
        return ""

    chunks = []
    for idx, item in enumerate(faculty, start=1):
        chunks.append(
            f"{idx}. {item['professor']} | "
            f"Research={item['research_fit_score']:.1f}/10 | "
            f"Activity={item['activity_score']:.1f}/10 | "
            f"Funding={item['funding_score']:.1f}/10 | "
            f"Overall={item['faculty_score']:.1f}/10 | "
            f"Area={', '.join(item['fields'][:3])} | "
            f"Paper={item['paper'][:100]} ({item['year']})"
        )
    return "\n".join(chunks)


def format_funding_details(faculty):
    if not faculty:
        return ""

    chunks = []
    for item in faculty:
        evidence = item["funding_evidence"]
        if evidence:
            chunks.append(f"{item['professor']}: {evidence}")
    return "\n".join(chunks)


def update_dataframe_row(df, idx, result):
    faculty = result["faculty"]

    if not faculty:
        df.at[idx, "Discovered_Field"] = ""
        df.at[idx, "General_PI"] = ""
        df.at[idx, "General_Paper"] = ""
        df.at[idx, "University_Research_Fit"] = 0
        df.at[idx, "University_Activity"] = 0
        df.at[idx, "University_Funding_Evidence"] = 0
        df.at[idx, "University_Score"] = 0
        df.at[idx, "Faculty_Breadth"] = 0
        df.at[idx, "Top_Faculty"] = ""
        df.at[idx, "Top_Faculty_Details"] = ""
        df.at[idx, "Funding_Evidence_Details"] = ""
        df.at[idx, "Last_Publication"] = ""
        return

    top_n = faculty[:MAX_FACULTY_PER_UNIVERSITY]

    research_score = sum(x["research_fit_score"] for x in top_n) / len(top_n)
    activity_score_avg = sum(x["activity_score"] for x in top_n) / len(top_n)
    funding_score_avg = sum(x["funding_score"] for x in top_n) / len(top_n)

    university_score = (
        0.60 * research_score
        + 0.20 * activity_score_avg
        + 0.20 * funding_score_avg
    )

    breadth = sum(1 for x in top_n if x["research_fit_score"] >= 6.5)

    last_pub = max(
        [x["latest_publication"] for x in top_n if x["latest_publication"]],
        default="",
    )

    first = top_n[0]
    df.at[idx, "Discovered_Field"] = first["field"]
    df.at[idx, "General_PI"] = first["professor"]
    df.at[idx, "General_Paper"] = f"{first['paper']} ({first['year']})"
    df.at[idx, "University_Research_Fit"] = round(research_score, 2)
    df.at[idx, "University_Activity"] = round(activity_score_avg, 2)
    df.at[idx, "University_Funding_Evidence"] = round(funding_score_avg, 2)
    df.at[idx, "University_Score"] = round(university_score, 2)
    df.at[idx, "Faculty_Breadth"] = breadth
    df.at[idx, "Top_Faculty"] = "; ".join(x["professor"] for x in top_n)
    df.at[idx, "Top_Faculty_Details"] = format_faculty_details(top_n)
    df.at[idx, "Funding_Evidence_Details"] = format_funding_details(top_n)
    df.at[idx, "Last_Publication"] = last_pub
    df.at[idx, "Note"] = first["profile_url"]


def style_excel_workbook(file_name, main_sheet_name="Sheet1"):
    """Apply practical formatting so each faculty is readable row-by-row."""
    try:
        wb = load_workbook(file_name)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.sheet_view.showGridLines = False

            # Sensible widths; cap very long text columns.
            for col_cells in ws.columns:
                values = [str(c.value or "") for c in col_cells[:50]]
                width = min(max(max((len(v) for v in values), default=10) + 2, 12), 42)
                ws.column_dimensions[col_cells[0].column_letter].width = width

            # Extra room for detailed narrative cells.
            for col in ws.iter_cols(1, ws.max_column):
                header = str(col[0].value or "")
                if "Details" in header or "Evidence" in header:
                    ws.column_dimensions[col[0].column_letter].width = 55
                    for cell in col[1:]:
                        cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

        wb.save(file_name)
    except Exception as exc:
        print(f"[!] Excel styling skipped: {exc}")


def build_faculty_rows(processed_results):
    """One faculty member per row for clean Excel filtering/sorting."""
    rows = []
    for result in processed_results:
        institution = result.get("institution") or {}
        uni = result.get("university", "")
        faculty = result.get("faculty", [])

        if not faculty:
            rows.append({
                "University": uni,
                "Faculty_Rank": "",
                "Professor": "",
                "University_OpenAlex": institution.get("display_name", ""),
                "Research_Fit": 0,
                "Activity": 0,
                "Funding_Evidence": 0,
                "Overall_Score": 0,
                "Research_Areas": "",
                "Matched_Works": 0,
                "Recent_Works": 0,
                "Latest_Publication": "",
                "Best_Paper": "",
                "Paper_Year": "",
                "Funding_Evidence_Details": "",
                "OpenAlex_Profile": "",
                "Homepage": "",
            })
            continue

        for rank, item in enumerate(faculty, start=1):
            rows.append({
                "University": uni,
                "Faculty_Rank": rank,
                "Professor": item.get("professor", ""),
                "University_OpenAlex": institution.get("display_name", ""),
                "Research_Fit": item.get("research_fit_score", 0),
                "Activity": item.get("activity_score", 0),
                "Funding_Evidence": item.get("funding_score", 0),
                "Overall_Score": item.get("faculty_score", 0),
                "Research_Areas": "; ".join(item.get("fields", [])[:3]),
                "Matched_Works": item.get("matched_works_count", 0),
                "Recent_Works": item.get("recent_works_count", 0),
                "Latest_Publication": item.get("latest_publication", ""),
                "Best_Paper": item.get("paper", ""),
                "Paper_Year": item.get("year", ""),
                "Funding_Evidence_Details": item.get("funding_evidence", ""),
                "OpenAlex_Profile": item.get("profile_url", ""),
                "Homepage": item.get("homepage_url", ""),
            })
    return pd.DataFrame(rows)


def save_outputs(df, processed_results):
    """Save the university tracker plus a row-by-row Faculty Results sheet."""
    faculty_df = build_faculty_rows(processed_results)

    if FILE_NAME.endswith(".csv"):
        df.to_csv(FILE_NAME, index=False)
        faculty_file = "Faculty Results.csv"
        faculty_df.to_csv(faculty_file, index=False)
        return faculty_file

    # Preserve the existing workbook and all unrelated sheets.
    wb = load_workbook(FILE_NAME)
    main_sheet_name = wb.sheetnames[0] if wb.sheetnames else "Sheet1"
    wb.close()

    with pd.ExcelWriter(
        FILE_NAME,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        df.to_excel(writer, sheet_name=main_sheet_name, index=False)
        faculty_df.to_excel(writer, sheet_name="Faculty Results", index=False)

    style_excel_workbook(FILE_NAME, main_sheet_name)
    return FILE_NAME


# ==============================================================================
# TELEGRAM REPORT
# ==============================================================================

def build_university_report(universities, all_matches):
    """Build multiple short Telegram messages; never silently truncate the report."""
    chunks = []
    current = ["🎯 <b>Faculty Finder v3 — Faculty Matches</b>", ""]

    def flush():
        nonlocal current
        if current and len("\n".join(current)) > 0:
            chunks.append("\n".join(current))
        current = []

    results_by_key = {university_key(item.get("university", "")): item for item in all_matches}

    for uni_name in universities:
        uni_result = results_by_key.get(university_key(uni_name), {"university": uni_name, "faculty": []})
        faculty = uni_result.get("faculty", [])

        lines = [f"🏛 <b>{html.escape(uni_name)}</b>"]
        if not faculty:
            lines.append("   No strong faculty match found.")
            lines.append("─" * 28)
        else:
            top_n = faculty[:MAX_FACULTY_PER_UNIVERSITY]
            research = sum(x["research_fit_score"] for x in top_n) / len(top_n)
            activity = sum(x["activity_score"] for x in top_n) / len(top_n)
            funding = sum(x["funding_score"] for x in top_n) / len(top_n)
            university_score = 0.60 * research + 0.20 * activity + 0.20 * funding

            lines.extend([
                f"   ⭐ University score: <b>{university_score:.1f}/10</b>",
                f"   🎯 Research: {research:.1f} | 📚 Activity: {activity:.1f} | 💰 Funding evidence: {funding:.1f}",
                "",
            ])

            for rank, item in enumerate(top_n, start=1):
                professor = html.escape(item.get("professor", ""))
                fields = html.escape(", ".join(item.get("fields", [])[:2]))
                paper = html.escape(item.get("paper", "")[:85])
                evidence = html.escape(item.get("funding_evidence", "No explicit evidence")[:110])
                profile = html.escape(item.get("profile_url", ""))

                lines.extend([
                    f"{rank}. <b>{professor}</b>",
                    f"   Fit: {item['research_fit_score']:.1f} | Activity: {item['activity_score']:.1f} | Funding: {item['funding_score']:.1f} | Overall: <b>{item['faculty_score']:.1f}/10</b>",
                    f"   Area: {fields}",
                    f"   Paper: <i>{paper}</i> ({item.get('year', '')})",
                    f"   Funding: {evidence}",
                    f"   <a href=\"{profile}\">OpenAlex profile</a>",
                    "",
                ])

            lines.append("─" * 28)

        block = "\n".join(lines)
        # Keep a conservative margin below Telegram's 4096-char limit.
        if len("\n".join(current + [block])) > 3500:
            flush()
            current = ["🎯 <b>Faculty Finder v3 — continued</b>", "", block]
        else:
            current.append(block)

    flush()
    return chunks


def send_report_chunks(chunks):
    for chunk in chunks:
        send_telegram_html(chunk)
        time.sleep(0.4)


# ==============================================================================
# MAIN
# ==============================================================================

def main():
    print("🚀 Running Faculty Finder v3...")
    print(f"Loading file: {FILE_NAME}")

    if not os.path.exists(FILE_NAME):
        print(f"[!] Error: {FILE_NAME} not found.")
        return

    try:
        if FILE_NAME.endswith(".csv"):
            df = pd.read_csv(FILE_NAME)
            original_sheet_name = None
        else:
            df = pd.read_excel(FILE_NAME)
            wb = load_workbook(FILE_NAME, read_only=True)
            original_sheet_name = wb.sheetnames[0] if wb.sheetnames else "Sheet1"
            wb.close()
    except Exception as exc:
        print(f"[!] Failed to read tracking file: {exc}")
        return

    if "University" not in df.columns:
        print("[!] Required column 'University' not found.")
        return

    df = ensure_output_columns(df)

    # Only process unchecked rows, but collapse duplicates to UNIQUE universities.
    mask = (
        df["University"].notna()
        & (df["University"].astype(str).str.strip() != "")
        & (df["University"].astype(str).str.lower() != "nan")
        & (df["Checked_General?"].astype(str).str.upper() != "YES")
    )

    seen_keys = set()
    unique_batch = []
    for idx in df.index[mask]:
        raw_name = str(df.at[idx, "University"]).strip()
        key = university_key(raw_name)
        if not key or key in seen_keys:
            continue
        seen_keys.add(key)
        unique_batch.append((key, raw_name))
        if len(unique_batch) >= BATCH_SIZE:
            break

    if not unique_batch:
        print("[✓] All unique universities have been checked.")
        send_telegram_html("🎉 <b>Faculty Finder v3:</b> all unique universities have been processed.")
        return

    print(f"\nProcessing {len(unique_batch)} UNIQUE universities (duplicates collapsed):")
    for number, (_, uni) in enumerate(unique_batch, start=1):
        duplicate_count = len(duplicate_indices_for_university(df, university_key(uni)))
        suffix = f" ({duplicate_count} rows in sheet)" if duplicate_count > 1 else ""
        print(f"  {number}. {uni}{suffix}")

    all_matches = []
    successful_count = 0
    failed_count = 0

    for batch_pos, (uni_key, raw_uni) in enumerate(unique_batch, start=1):
        print("\n" + "=" * 88)
        print(f"🏛 [{batch_pos}/{len(unique_batch)}] {raw_uni}")
        print("=" * 88)

        result = search_faculty_broad(raw_uni)

        if not result["ok"]:
            print(f"[!] Search failed for {raw_uni}; leaving its rows unchecked.")
            failed_count += 1
            continue

        successful_count += 1
        all_matches.append({
            "university": raw_uni,
            "institution": result.get("institution") or {},
            "faculty": result.get("faculty", []),
        })

        # IMPORTANT: update ALL duplicate rows for this university in one pass.
        matching_indices = duplicate_indices_for_university(df, uni_key)
        for idx in matching_indices:
            update_dataframe_row(df, idx, result)
            df.at[idx, "Checked_General?"] = "YES"

        print(f"[✓] Updated {len(matching_indices)} spreadsheet row(s) for {raw_uni}.")

        if result["faculty"]:
            print(f"[✓] Top {min(len(result['faculty']), MAX_FACULTY_PER_UNIVERSITY)} faculty:")
            for rank, faculty in enumerate(result["faculty"], start=1):
                print(
                    f"   {rank}. {faculty['professor']} | "
                    f"Fit={faculty['research_fit_score']:.1f} | "
                    f"Activity={faculty['activity_score']:.1f} | "
                    f"Funding={faculty['funding_score']:.1f} | "
                    f"Overall={faculty['faculty_score']:.1f} | "
                    f"{faculty['field']}"
                )
        else:
            print("[–] No strong faculty match found.")

        time.sleep(0.5)

    # Save progress even if some universities failed.
    try:
        faculty_file = save_outputs(df, all_matches)
        print(f"\n[✓] Tracking output updated: {FILE_NAME}")
        if faculty_file != FILE_NAME:
            print(f"[✓] Row-by-row faculty output: {faculty_file}")
    except Exception as exc:
        print(f"[!] Failed to save output files: {exc}")
        return

    # Telegram: multiple messages so no university/faculty block is silently cut.
    universities = [uni for _, uni in unique_batch]
    report_chunks = build_university_report(universities, all_matches)
    send_report_chunks(report_chunks)

    send_telegram_document(
        FILE_NAME,
        caption=(
            f"💾 Faculty Finder v3 updated sheet | "
            f"Processed: {successful_count}/{len(unique_batch)} unique universities | "
            f"Failed: {failed_count}"
        ),
    )

    if FILE_NAME.endswith(".csv") and os.path.exists("Faculty Results.csv"):
        send_telegram_document(
            "Faculty Results.csv",
            caption="📋 Row-by-row faculty results (one professor per row).",
        )

    print("\n✅ Run complete.")

if __name__ == "__main__":
    main()
